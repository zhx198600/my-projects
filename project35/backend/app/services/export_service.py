from typing import List, Dict, Any, Optional
from io import BytesIO
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


class ExportService:
    @staticmethod
    def _get_timestamp() -> str:
        return datetime.now().strftime("%Y%m%d_%H%M%S")

    @staticmethod
    def generate_filename(report_name: str, extension: str) -> str:
        timestamp = ExportService._get_timestamp()
        return f"{report_name}_{timestamp}.{extension}"

    @staticmethod
    def _style_excel_header(ws, columns: List[str]):
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        for col_num, column_title in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = column_title
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

    @staticmethod
    def _style_excel_data(ws, row_count: int, col_count: int):
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for row in range(1, row_count + 1):
            for col in range(1, col_count + 1):
                cell = ws.cell(row=row, column=col)
                cell.border = thin_border
                if row > 1:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

    @staticmethod
    def _auto_adjust_column_width(ws, col_count: int):
        for col in range(1, col_count + 1):
            max_length = 0
            column_letter = get_column_letter(col)
            for cell in ws[column_letter]:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    @staticmethod
    def export_budget_summary_to_excel(data: List[Dict[str, Any]]) -> BytesIO:
        output = BytesIO()
        wb = Workbook()
        ws = wb.active
        ws.title = "预算汇总"

        columns = [
            "模板名称", "预算期间", "部门", "科目", "月份", 
            "汇总类型", "版本", "预算金额", "已使用", "已占用"
        ]
        
        ExportService._style_excel_header(ws, columns)

        summary_type_map = {
            "detail": "明细",
            "department_total": "部门汇总",
            "subject_total": "科目汇总",
            "grand_total": "总计"
        }

        for row_num, item in enumerate(data, 2):
            ws.cell(row=row_num, column=1, value=item.get("template_name", ""))
            ws.cell(row=row_num, column=2, value=item.get("period_name", ""))
            ws.cell(row=row_num, column=3, value=item.get("department_name", "-"))
            ws.cell(row=row_num, column=4, value=item.get("subject_name", "-"))
            ws.cell(row=row_num, column=5, value=f"{item.get('month', '')}月" if item.get("month") else "-")
            ws.cell(row=row_num, column=6, value=summary_type_map.get(item.get("summary_type", ""), item.get("summary_type", "")))
            ws.cell(row=row_num, column=7, value=item.get("version", ""))
            ws.cell(row=row_num, column=8, value=float(item.get("budget_amount", 0)))
            ws.cell(row=row_num, column=9, value=float(item.get("used_amount", 0)))
            ws.cell(row=row_num, column=10, value=float(item.get("occupied_amount", 0)))

        row_count = len(data) + 1
        col_count = len(columns)
        
        ExportService._style_excel_data(ws, row_count, col_count)
        ExportService._auto_adjust_column_width(ws, col_count)

        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_budget_execution_to_excel(data: List[Dict[str, Any]]) -> BytesIO:
        output = BytesIO()
        wb = Workbook()
        ws = wb.active
        ws.title = "预算执行"

        columns = [
            "预算期间", "部门", "科目", "预算金额", "已发生", 
            "已占用", "剩余金额", "执行率(%)"
        ]
        
        ExportService._style_excel_header(ws, columns)

        for row_num, item in enumerate(data, 2):
            ws.cell(row=row_num, column=1, value=item.get("period_name", ""))
            ws.cell(row=row_num, column=2, value=item.get("department_name", "-"))
            ws.cell(row=row_num, column=3, value=item.get("subject_name", "-"))
            ws.cell(row=row_num, column=4, value=float(item.get("budget_amount", 0)))
            ws.cell(row=row_num, column=5, value=float(item.get("used_amount", 0)))
            ws.cell(row=row_num, column=6, value=float(item.get("occupied_amount", 0)))
            ws.cell(row=row_num, column=7, value=float(item.get("remaining_amount", 0)))
            ws.cell(row=row_num, column=8, value=float(item.get("execution_rate", 0)))

        row_count = len(data) + 1
        col_count = len(columns)
        
        ExportService._style_excel_data(ws, row_count, col_count)
        ExportService._auto_adjust_column_width(ws, col_count)

        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def _get_pdf_styles() -> Dict[str, ParagraphStyle]:
        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=16,
            alignment=1,
            spaceAfter=20,
            textColor=colors.HexColor("#1890ff")
        )
        
        header_style = ParagraphStyle(
            'CustomHeader',
            parent=styles['Normal'],
            fontSize=10,
            alignment=1,
            textColor=colors.white,
            fontName='Helvetica-Bold'
        )
        
        normal_style = ParagraphStyle(
            'CustomNormal',
            parent=styles['Normal'],
            fontSize=9,
            alignment=0,
            leading=12
        )
        
        return {
            'title': title_style,
            'header': header_style,
            'normal': normal_style
        }

    @staticmethod
    def export_budget_summary_to_pdf(data: List[Dict[str, Any]], title: str = "预算汇总报表") -> BytesIO:
        output = BytesIO()
        doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        
        styles = ExportService._get_pdf_styles()
        elements = []

        elements.append(Paragraph(title, styles['title']))
        elements.append(Spacer(1, 12))

        summary_type_map = {
            "detail": "明细",
            "department_total": "部门汇总",
            "subject_total": "科目汇总",
            "grand_total": "总计"
        }

        table_data = [
            ["模板名称", "预算期间", "部门", "科目", "月份", "汇总类型", "版本", "预算金额", "已使用", "已占用"]
        ]

        for item in data:
            row = [
                str(item.get("template_name", "")),
                str(item.get("period_name", "")),
                str(item.get("department_name", "-")),
                str(item.get("subject_name", "-")),
                f"{item.get('month', '')}月" if item.get("month") else "-",
                summary_type_map.get(item.get("summary_type", ""), item.get("summary_type", "")),
                str(item.get("version", "")),
                f"¥{float(item.get('budget_amount', 0)):,.2f}",
                f"¥{float(item.get('used_amount', 0)):,.2f}",
                f"¥{float(item.get('occupied_amount', 0)):,.2f}"
            ]
            table_data.append(row)

        col_widths = [1.2 * inch, 1 * inch, 1 * inch, 1 * inch, 0.7 * inch, 1 * inch, 0.7 * inch, 1.2 * inch, 1.2 * inch, 1.2 * inch]

        table = Table(table_data, colWidths=col_widths)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#4472C4")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('GRID', (0, 0), (-1, -1), 1, colors.gray),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f5f5")]),
        ]))

        elements.append(table)
        
        doc.build(elements)
        output.seek(0)
        return output

    @staticmethod
    def export_budget_execution_to_pdf(data: List[Dict[str, Any]], title: str = "预算执行报表") -> BytesIO:
        output = BytesIO()
        doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        
        styles = ExportService._get_pdf_styles()
        elements = []

        elements.append(Paragraph(title, styles['title']))
        elements.append(Spacer(1, 12))

        table_data = [
            ["预算期间", "部门", "科目", "预算金额", "已发生", "已占用", "剩余金额", "执行率"]
        ]

        for item in data:
            row = [
                str(item.get("period_name", "")),
                str(item.get("department_name", "-")),
                str(item.get("subject_name", "-")),
                f"¥{float(item.get('budget_amount', 0)):,.2f}",
                f"¥{float(item.get('used_amount', 0)):,.2f}",
                f"¥{float(item.get('occupied_amount', 0)):,.2f}",
                f"¥{float(item.get('remaining_amount', 0)):,.2f}",
                f"{float(item.get('execution_rate', 0)):.2f}%"
            ]
            table_data.append(row)

        col_widths = [1.2 * inch, 1.2 * inch, 1.2 * inch, 1.3 * inch, 1.3 * inch, 1.3 * inch, 1.3 * inch, 1 * inch]

        table = Table(table_data, colWidths=col_widths)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#52c41a")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('GRID', (0, 0), (-1, -1), 1, colors.gray),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f5f5")]),
        ]))

        elements.append(table)
        
        doc.build(elements)
        output.seek(0)
        return output

    @staticmethod
    def batch_export_to_excel(reports: List[Dict[str, Any]]) -> BytesIO:
        output = BytesIO()
        wb = Workbook()
        
        if wb.active:
            wb.remove(wb.active)

        for report in reports:
            report_type = report.get("type", "")
            report_data = report.get("data", [])
            sheet_name = report.get("sheet_name", report_type)

            ws = wb.create_sheet(title=sheet_name[:31])

            if report_type == "budget_summary":
                columns = [
                    "模板名称", "预算期间", "部门", "科目", "月份", 
                    "汇总类型", "版本", "预算金额", "已使用", "已占用"
                ]
                summary_type_map = {
                    "detail": "明细",
                    "department_total": "部门汇总",
                    "subject_total": "科目汇总",
                    "grand_total": "总计"
                }
                ExportService._style_excel_header(ws, columns)
                for row_num, item in enumerate(report_data, 2):
                    ws.cell(row=row_num, column=1, value=item.get("template_name", ""))
                    ws.cell(row=row_num, column=2, value=item.get("period_name", ""))
                    ws.cell(row=row_num, column=3, value=item.get("department_name", "-"))
                    ws.cell(row=row_num, column=4, value=item.get("subject_name", "-"))
                    ws.cell(row=row_num, column=5, value=f"{item.get('month', '')}月" if item.get("month") else "-")
                    ws.cell(row=row_num, column=6, value=summary_type_map.get(item.get("summary_type", ""), item.get("summary_type", "")))
                    ws.cell(row=row_num, column=7, value=item.get("version", ""))
                    ws.cell(row=row_num, column=8, value=float(item.get("budget_amount", 0)))
                    ws.cell(row=row_num, column=9, value=float(item.get("used_amount", 0)))
                    ws.cell(row=row_num, column=10, value=float(item.get("occupied_amount", 0)))
                
                ExportService._style_excel_data(ws, len(report_data) + 1, len(columns))
                ExportService._auto_adjust_column_width(ws, len(columns))

            elif report_type == "budget_execution":
                columns = [
                    "预算期间", "部门", "科目", "预算金额", "已发生", 
                    "已占用", "剩余金额", "执行率(%)"
                ]
                ExportService._style_excel_header(ws, columns)
                for row_num, item in enumerate(report_data, 2):
                    ws.cell(row=row_num, column=1, value=item.get("period_name", ""))
                    ws.cell(row=row_num, column=2, value=item.get("department_name", "-"))
                    ws.cell(row=row_num, column=3, value=item.get("subject_name", "-"))
                    ws.cell(row=row_num, column=4, value=float(item.get("budget_amount", 0)))
                    ws.cell(row=row_num, column=5, value=float(item.get("used_amount", 0)))
                    ws.cell(row=row_num, column=6, value=float(item.get("occupied_amount", 0)))
                    ws.cell(row=row_num, column=7, value=float(item.get("remaining_amount", 0)))
                    ws.cell(row=row_num, column=8, value=float(item.get("execution_rate", 0)))
                
                ExportService._style_excel_data(ws, len(report_data) + 1, len(columns))
                ExportService._auto_adjust_column_width(ws, len(columns))

        wb.save(output)
        output.seek(0)
        return output
